from flask import Flask, request, jsonify
from werkzeug.utils import secure_filename
import os
import google.generativeai as genai
from dotenv import load_dotenv
import PyPDF2
from docx import Document
import json
import re

# Load environment variables
load_dotenv()

app = Flask(__name__)

UPLOAD_FOLDER = './uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max file size

# Configure Gemini API
API_KEY = os.getenv('GOOGLE_API_KEY')
if API_KEY:
    genai.configure(api_key=API_KEY)
    
    safety_settings = {
        "HARM_CATEGORY_HARASSMENT": "block_none",
        "HARM_CATEGORY_HATE_SPEECH": "block_none",
        "HARM_CATEGORY_SEXUALLY_EXPLICIT": "block_none",
        "HARM_CATEGORY_DANGEROUS_CONTENT": "block_none",
    }
    
    model = genai.GenerativeModel(
        "models/gemini-2.5-flash",
        safety_settings=safety_settings
    )
else:
    print("WARNING: GOOGLE_API_KEY not found in environment variables")
    model = None

# Allowed file extensions
ALLOWED_EXTENSIONS = {'pdf', 'txt', 'docx'}

# --- PROMPTS ---
DETAILED_CLAUSE_ANALYSIS_PROMPT = """
You are a legal expert analyzing contracts and agreements for LexiGuard. 
Analyze the following document and identify ALL risky or concerning clauses with deep explanations.

For EACH risky clause, provide:
1. **clause**: The exact clause text or relevant excerpt (keep it under 200 characters if too long)
2. **risk_level**: "High", "Medium", or "Low"
3. **impact**: Brief description of potential harm to the user (1 sentence)
4. **recommendation**: Specific actionable advice for negotiation (1-2 sentences)
5. **explanation**: Detailed plain-language explanation of why this clause is risky and what could go wrong (2-3 sentences)

Focus on these risk categories:
- Termination rights (sudden eviction, firing without cause)
- Financial liability (unlimited damages, penalties)
- Automatic renewal traps
- Non-compete restrictions
- Indemnification clauses
- Limitation of liability
- Unilateral changes by one party
- Waiver of legal rights

You MUST respond ONLY with a valid JSON array. Use this exact format:
[
    {
        "clause": "Original clause text here",
        "risk_level": "High",
        "impact": "Brief impact description",
        "recommendation": "What the user should do or negotiate",
        "explanation": "Detailed plain-language explanation of why this is risky"
    }
]

If no risks are found, return an empty array: []
CRITICAL: Respond ONLY with valid JSON, no additional text, no markdown.
"""

DOCUMENT_EMAIL_PROMPT = """
You are LexiGuard, an AI assistant that helps users communicate about legal document reviews.

Generate a professional email to send to a legal advisor, counterparty, or stakeholder regarding a legal document review.

Document Summary:
{document_summary}

Identified Risks:
{risk_summary}

The email should:
1. Be professional, clear, and concise
2. Summarize the key findings from the document analysis
3. Highlight the most critical risks identified
4. Request clarification, revision, or discussion on the concerning clauses
5. Maintain a collaborative and constructive tone
6. Be ready to send with minimal editing

Generate ONLY the email body. Do not include:
- Subject line
- Sender/recipient names or addresses
- Placeholders like [Your Name] or [Company Name]

Start directly with a professional greeting and the email content.
"""

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_pdf(filepath):
    """Extract text content from a PDF file"""
    try:
        with open(filepath, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            return text
    except Exception as e:
        raise Exception(f"Error extracting PDF text: {str(e)}")

def extract_text_from_docx(filepath):
    """Extract text content from a DOCX file"""
    try:
        doc = Document(filepath)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text
    except Exception as e:
        raise Exception(f"Error extracting DOCX text: {str(e)}")

def analyze_document_with_gemini(text, analysis_type="general"):
    """
    Analyze document text using Gemini AI
    analysis_type: 'general' or 'detailed_clauses'
    """
    if not model:
        raise Exception("Gemini API not configured. Please set GOOGLE_API_KEY environment variable.")
    
    if analysis_type == "general":
        prompt = f"""
        Analyze the following legal document and provide:
        1. A brief summary (2-3 sentences)
        2. A list of potential risks or concerning clauses
        
        Document text:
        {text[:8000]}
        
        Respond in JSON format:
        {{
            "summary": "Brief summary here",
            "risks": ["Risk 1", "Risk 2", ...]
        }}
        """
    else:  # detailed_clauses
        prompt = f"{DETAILED_CLAUSE_ANALYSIS_PROMPT}\n\nDocument:\n{text[:12000]}"
    
    try:
        response = model.generate_content(prompt)
        response_text = response.text.strip()
        
        # Clean up response to extract JSON
        response_text = re.sub(r'^```json\s*', '', response_text)
        response_text = re.sub(r'^```\s*', '', response_text)
        response_text = re.sub(r'\s*```$', '', response_text)
        response_text = response_text.strip()
        
        # Parse JSON
        result = json.loads(response_text)
        return result
    
    except json.JSONDecodeError as e:
        print(f"JSON Parse Error: {e}")
        print(f"Raw response: {response_text[:500]}")
        # Return fallback structure
        if analysis_type == "general":
            return {
                "summary": "Analysis completed but formatting error occurred.",
                "risks": ["Unable to parse detailed risks. Please try again."]
            }
        else:
            return []
    
    except Exception as e:
        raise Exception(f"Error during Gemini analysis: {str(e)}")

# --- ROUTES ---

@app.route('/analyze-pdf', methods=['POST'])
def analyze_pdf():
    """
    Legacy endpoint - General document analysis
    """
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    
    file = request.files['file']
    
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    if not allowed_file(file.filename):
        return jsonify({'error': 'Invalid file type. Only PDF, TXT, and DOCX allowed'}), 400
    
    try:
        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        
        # Extract text based on file type
        if filename.endswith('.pdf'):
            document_text = extract_text_from_pdf(filepath)
        elif filename.endswith('.docx'):
            document_text = extract_text_from_docx(filepath)
        else:
            with open(filepath, 'r', encoding='utf-8') as f:
                document_text = f.read()
        
        # Analyze with Gemini (general analysis)
        analysis_result = analyze_document_with_gemini(document_text, analysis_type="general")
        
        # Add filename to response
        analysis_result['filename'] = filename
        analysis_result['text_preview'] = document_text[:500]
        
        return jsonify(analysis_result)
    
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/analyze-clauses', methods=['POST'])
def analyze_clauses():
    """
    NEW ENDPOINT - Detailed clause-by-clause risk analysis with explanations
    
    Example Response:
    {
        "filename": "lease_agreement.pdf",
        "total_risky_clauses": 3,
        "clauses": [
            {
                "clause": "Landlord may terminate with 7 days notice",
                "risk_level": "High",
                "impact": "Could lead to sudden displacement without adequate time to find housing",
                "recommendation": "Request amendment to 30-day notice period minimum",
                "explanation": "This clause allows your landlord to terminate the lease with one week's notice. This could lead to sudden displacement — consider requesting a 30-day notice period."
            }
        ],
        "document_preview": "..."
    }
    """
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    
    file = request.files['file']
    
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    if not allowed_file(file.filename):
        return jsonify({'error': 'Invalid file type. Only PDF, TXT, and DOCX allowed'}), 400
    
    try:
        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        
        # Extract text based on file type
        if filename.endswith('.pdf'):
            document_text = extract_text_from_pdf(filepath)
        elif filename.endswith('.docx'):
            document_text = extract_text_from_docx(filepath)
        else:
            with open(filepath, 'r', encoding='utf-8') as f:
                document_text = f.read()
        
        # Analyze with Gemini (detailed clause analysis)
        clauses = analyze_document_with_gemini(document_text, analysis_type="detailed_clauses")
        
        # Ensure we always return a list
        if not isinstance(clauses, list):
            clauses = []
        
        return jsonify({
            'filename': filename,
            'total_risky_clauses': len(clauses),
            'clauses': clauses,
            'document_preview': document_text[:300]
        })
    
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/analyze', methods=['POST'])
def analyze_text():
    """
    Analyze raw text input (no file upload)
    Supports both 'general' and 'detailed' analysis types
    """
    data = request.get_json()
    text = data.get('text', '')
    analysis_type = data.get('type', 'general')  # 'general' or 'detailed'
    
    if not text:
        return jsonify({'error': 'No text provided'}), 400
    
    try:
        if analysis_type == 'detailed':
            result = analyze_document_with_gemini(text, analysis_type="detailed_clauses")
            return jsonify({
                'total_risky_clauses': len(result) if isinstance(result, list) else 0,
                'clauses': result if isinstance(result, list) else []
            })
        else:
            result = analyze_document_with_gemini(text, analysis_type="general")
            return jsonify(result)
    
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/draft-negotiation', methods=['POST'])
def draft_negotiation():
    """
    Generate a negotiation email for a specific risky clause.
    
    Expected JSON body:
    {
        "clause": "The specific clause text to negotiate"
    }
    """
    data = request.get_json()
    
    if not data:
        return jsonify({'error': 'No data provided'}), 400
    
    clause = data.get('clause', '')
    
    if not clause:
        return jsonify({'error': 'Clause text is required'}), 400
    
    if not model:
        return jsonify({'error': 'Gemini API not configured'}), 500
    
    try:
        prompt = f"""
You are LexiGuard, an AI that helps users negotiate risky clauses politely.
Draft a short, professional, and polite email asking to revise or clarify the following clause:

Clause:
{clause}

Do not include placeholders like [Company]. Just write a natural, ready-to-send email body.
"""
        
        response = model.generate_content(prompt)
        email_text = response.text.strip() or "Could not generate negotiation email."
        
        return jsonify({'negotiation_email': email_text})
    
    except Exception as e:
        return jsonify({'error': f'Error generating negotiation email: {str(e)}'}), 500

@app.route('/draft-document-email', methods=['POST'])
def draft_document_email():
    """
    Generate a comprehensive professional email covering all document findings.
    Supports both standard and detailed analysis results.
    
    Expected JSON body:
    {
        "document_summary": "Summary or preview of the document",
        "risk_summary": "Summary of identified risks"
    }
    """
    data = request.get_json()
    
    if not data:
        return jsonify({'error': 'No data provided'}), 400
    
    document_summary = data.get('document_summary', '')
    risk_summary = data.get('risk_summary', '')
    
    if not document_summary or not risk_summary:
        return jsonify({'error': 'Both document_summary and risk_summary are required'}), 400
    
    if not model:
        return jsonify({'error': 'Gemini API not configured'}), 500
    
    try:
        prompt = DOCUMENT_EMAIL_PROMPT.format(
            document_summary=document_summary[:2000],
            risk_summary=risk_summary[:2000]
        )
        
        response = model.generate_content(prompt)
        email_text = response.text.strip() or "Could not generate email."
        
        return jsonify({'document_email': email_text})
    
    except Exception as e:
        return jsonify({'error': f'Error generating email: {str(e)}'}), 500

@app.route('/chat', methods=['POST'])
def chat_with_document():
    """
    Chat with document context.
    
    Expected JSON body:
    {
        "message": "User's question",
        "document_text": "The document content for context"
    }
    """
    data = request.get_json()
    
    if not data:
        return jsonify({'error': 'No data provided'}), 400
    
    message = data.get('message', '')
    document_text = data.get('document_text', '')
    
    if not message or not document_text:
        return jsonify({'error': 'Both message and document_text are required'}), 400
    
    if not model:
        return jsonify({'error': 'Gemini API not configured'}), 500
    
    try:
        prompt = f"""
You are LexiGuard, a helpful AI assistant. Answer the user's messages based only on the provided legal document.
Maintain conversation context. Do not assume beyond the document.

Document:
{document_text[:8000]}

User Message:
{message}

Respond concisely and clearly for a non-lawyer.
"""
        
        response = model.generate_content(prompt)
        answer = response.text.strip() or "No answer could be generated."
        
        return jsonify({'reply': answer})
    
    except Exception as e:
        return jsonify({'error': f'Error during chat: {str(e)}'}), 500

@app.route('/health', methods=['GET'])
def health_check():
    """Health check endpoint"""
    return jsonify({
        'status': 'healthy',
        'version': '1.3.0',
        'gemini_configured': model is not None
    })

@app.route('/', methods=['GET'])
def root():
    """Root endpoint with API information"""
    return jsonify({
        'message': 'LexiGuard Flask API is running',
        'version': '1.3.0',
        'endpoints': {
            'general_analysis': '/analyze-pdf (POST)',
            'text_analysis': '/analyze (POST)',
            'detailed_clause_analysis': '/analyze-clauses (POST)',
            'draft_negotiation': '/draft-negotiation (POST)',
            'draft_document_email': '/draft-document-email (POST)',
            'chat': '/chat (POST)',
            'health': '/health (GET)'
        }
    })

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8080))  # Google Cloud Run sets PORT to 8080
    print(f"Starting Flask server on host=0.0.0.0, port={port}")
    app.run(host="0.0.0.0", port=port, debug=False)  # debug=False for production